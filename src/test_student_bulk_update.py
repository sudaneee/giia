import io
from datetime import date

import openpyxl
from django.contrib.auth.models import User
from django.test import Client, TestCase
from django.urls import reverse

from src.models import Guardian, SchoolClass, Section, Student
from src.student_bulk_update import COLUMNS, SESSION_KEY
from wallet.test_support import seed_site_context_fixtures


def make_workbook(rows, headers=None):
    """rows: list of dicts keyed by Excel header; missing keys become blank cells."""
    headers = headers or [h for h, _ in COLUMNS]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = 'students.xlsx'
    return buf


class StudentBulkUpdateTests(TestCase):
    def setUp(self):
        seed_site_context_fixtures()
        section = Section.objects.create(name='Primary')
        self.class_a = SchoolClass.objects.create(name='Basic 1', level='B1', arm='A', section=section)
        self.class_b = SchoolClass.objects.create(name='Basic 2', level='B2', arm='A', section=section)
        self.staff = User.objects.create_user(username='staff', password='pw', is_staff=True)
        self.client = Client()
        self.client.force_login(self.staff)

        self.ali = Student.objects.create(
            admission_number='GIIA-1', first_name='Ali', last_name='Musa', enrolled_class=self.class_a,
            phone_number='08011111111', status='active', date_of_birth=date(2015, 1, 1),
        )
        self.zainab = Student.objects.create(
            admission_number='GIIA-2', first_name='Zainab', last_name='Umar', enrolled_class=self.class_a,
            status='active',
        )
        self.other = Student.objects.create(
            admission_number='GIIA-3', first_name='Other', last_name='Kid', enrolled_class=self.class_b,
        )

    def upload_preview(self, rows, class_obj=None):
        class_obj = class_obj or self.class_a
        return self.client.post(
            reverse('student_bulk_update'),
            {'class_id': class_obj.id, 'excel_file': make_workbook(rows)},
        )

    def test_export_downloads_only_that_class_with_all_columns(self):
        r = self.client.get(reverse('student_bulk_export'), {'class_id': self.class_a.id})
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb.active
        self.assertEqual([c.value for c in ws[1]], [h for h, _ in COLUMNS])
        admissions = sorted(ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1))
        self.assertEqual(admissions, ['GIIA-1', 'GIIA-2'])

    def test_partial_row_only_changes_the_filled_fields(self):
        self.upload_preview([{'Admission Number': 'GIIA-1', 'Status (active/inactive/graduated/suspended)': 'inactive'}])
        preview = self.client.session[SESSION_KEY]
        self.assertEqual(len(preview['rows']), 1)
        self.assertEqual(preview['rows'][0]['student_updates'], {'status': 'inactive'})

        self.client.post(reverse('student_bulk_update'), {'action': 'apply'})
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.status, 'inactive')
        self.assertEqual(self.ali.phone_number, '08011111111')  # untouched - blank in file
        self.assertEqual(self.ali.first_name, 'Ali')

    def test_preview_saves_nothing(self):
        self.upload_preview([{'Admission Number': 'GIIA-1', 'Status (active/inactive/graduated/suspended)': 'inactive'}])
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.status, 'active')

    def test_invalid_status_is_rejected_for_that_row_only(self):
        self.upload_preview([
            {'Admission Number': 'GIIA-1', 'Status (active/inactive/graduated/suspended)': 'ACTIVEX'},
            {'Admission Number': 'GIIA-2', 'Status (active/inactive/graduated/suspended)': 'INACTIVE'},
        ])
        preview = self.client.session[SESSION_KEY]
        self.assertEqual([r['admission_number'] for r in preview['rows']], ['GIIA-2'])
        errors = self.client.session['student_bulk_update_errors']
        self.assertTrue(any('GIIA-1' in e for e in errors))

    def test_student_from_another_class_is_rejected(self):
        self.upload_preview([{'Admission Number': 'GIIA-3', 'First Name': 'Changed'}])
        preview = self.client.session[SESSION_KEY]
        self.assertEqual(preview['rows'], [])
        self.assertTrue(any('not in' in e for e in self.client.session['student_bulk_update_errors']))

    def test_class_column_is_ignored_even_if_changed(self):
        self.upload_preview([{'Admission Number': 'GIIA-1', 'Class (reference only - not changed)': 'Basic 2 -- A'}])
        self.assertEqual(self.client.session.get(SESSION_KEY, {}).get('rows', []), [])
        self.client.post(reverse('student_bulk_update'), {'action': 'apply'})
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.enrolled_class, self.class_a)

    def test_guardian_is_updated_and_created_when_missing(self):
        self.upload_preview([{
            'Admission Number': 'GIIA-2',
            'Guardian First Name': 'Umar', 'Guardian Last Name': 'Bello', 'Guardian Phone Number': '08022222222',
        }])
        self.client.post(reverse('student_bulk_update'), {'action': 'apply'})
        guardian = self.zainab.guardians.first()
        self.assertIsNotNone(guardian)
        self.assertEqual(guardian.first_name, 'Umar')
        self.assertEqual(guardian.phone_number, '08022222222')

    def test_cancel_changes_nothing(self):
        self.upload_preview([{'Admission Number': 'GIIA-1', 'Status (active/inactive/graduated/suspended)': 'inactive'}])
        self.client.post(reverse('student_bulk_update'), {'action': 'cancel'})
        self.ali.refresh_from_db()
        self.assertEqual(self.ali.status, 'active')
        self.assertNotIn(SESSION_KEY, self.client.session)

    def test_file_without_admission_number_column_is_refused(self):
        bad = make_workbook([{'First Name': 'X'}], headers=['First Name'])
        r = self.client.post(reverse('student_bulk_update'), {'class_id': self.class_a.id, 'excel_file': bad}, follow=True)
        self.assertNotIn(SESSION_KEY, self.client.session)
