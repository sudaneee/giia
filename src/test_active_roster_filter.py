import io
from datetime import date

import openpyxl
from django.contrib.auth.models import User
from django.test import Client, TestCase
from django.urls import reverse

from src.models import (
    MidTermResult, Result, Section, SchoolClass, Session, Student, Subject, Term,
    TahfeezResult,
)
from wallet.test_support import seed_site_context_fixtures
from website.models import Picture


class ActiveRosterFixtureMixin:
    def setUp(self):
        seed_site_context_fixtures()
        Picture.objects.get_or_create(title='logo', defaults={'image': 'pics/placeholder.png'})
        section = Section.objects.create(name='Primary')
        self.school_class = SchoolClass.objects.create(name='Basic 1', level='B1', arm='A', section=section)
        self.session = Session.objects.create(name='2025/2026', start_date=date(2025, 9, 1), end_date=date(2026, 7, 31))
        self.term = Term.objects.create(name='First Term', session=self.session, start_date=date(2025, 9, 1), end_date=date(2025, 12, 1))
        self.subject = Subject.objects.create(name='Maths')
        self.active_student = Student.objects.create(
            admission_number='GIIA-1', first_name='Ali', last_name='Musa',
            enrolled_class=self.school_class, status='active',
        )
        self.inactive_student = Student.objects.create(
            admission_number='GIIA-2', first_name='Sara', last_name='Bello',
            enrolled_class=self.school_class, status='inactive',
        )
        self.staff = User.objects.create_user(username='staff', password='pw', is_staff=True)
        self.client = Client()
        self.client.force_login(self.staff)


class ResultEntryActiveRosterTests(ActiveRosterFixtureMixin, TestCase):
    def test_default_excludes_inactive_student(self):
        response = self.client.post(reverse('result_entry'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
        })
        students = list(response.context['students'])
        self.assertIn(self.active_student, students)
        self.assertNotIn(self.inactive_student, students)

    def test_include_inactive_checkbox_includes_them(self):
        response = self.client.post(reverse('result_entry'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
            'include_inactive': 'on',
        })
        students = list(response.context['students'])
        self.assertIn(self.active_student, students)
        self.assertIn(self.inactive_student, students)


class ResultUpdateActiveRosterTests(ActiveRosterFixtureMixin, TestCase):
    def setUp(self):
        super().setUp()
        Result.objects.create(
            student=self.active_student, subject=self.subject, class_assigned=self.school_class,
            session=self.session, term=self.term,
            ca1_marks=5, ca2_marks=5, home_work_marks=5, activity_marks=5, exam_marks=10,
        )
        Result.objects.create(
            student=self.inactive_student, subject=self.subject, class_assigned=self.school_class,
            session=self.session, term=self.term,
            ca1_marks=5, ca2_marks=5, home_work_marks=5, activity_marks=5, exam_marks=10,
        )

    def test_default_excludes_inactive_student_result(self):
        response = self.client.post(reverse('result_update'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
        })
        result_students = [r.student for r in response.context['results']]
        self.assertIn(self.active_student, result_students)
        self.assertNotIn(self.inactive_student, result_students)

    def test_include_inactive_includes_their_result(self):
        response = self.client.post(reverse('result_update'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
            'include_inactive': 'on',
        })
        result_students = [r.student for r in response.context['results']]
        self.assertIn(self.inactive_student, result_students)


class DownloadTemplateActiveRosterTests(ActiveRosterFixtureMixin, TestCase):
    def _names_in_workbook(self, content):
        workbook = openpyxl.load_workbook(io.BytesIO(content))
        sheet = workbook.active
        return {row[1] for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True) if row[1]}

    def test_default_excludes_inactive_student(self):
        response = self.client.post(reverse('download_template'), {'class_assigned': self.school_class.id})
        names = self._names_in_workbook(response.content)
        self.assertIn('Ali Musa', names)
        self.assertNotIn('Sara Bello', names)

    def test_include_inactive_includes_them(self):
        response = self.client.post(reverse('download_template'), {
            'class_assigned': self.school_class.id, 'include_inactive': 'on',
        })
        names = self._names_in_workbook(response.content)
        self.assertIn('Sara Bello', names)


class UploadResultsActiveRosterTests(ActiveRosterFixtureMixin, TestCase):
    def _build_workbook(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(['Student ID', 'Student Name', '1st CA Marks', '2nd CA Marks', 'Home Work Marks', 'Activity Marks', 'Exam Marks', 'Admission Number'])
        for student in (self.active_student, self.inactive_student):
            sheet.append([student.id, f'{student.first_name} {student.last_name}', 5, 5, 5, 5, 10, student.admission_number])
        buf = io.BytesIO()
        workbook.save(buf)
        buf.seek(0)
        return buf

    def test_default_skips_inactive_student_row(self):
        upload_file = self._build_workbook()
        upload_file.name = 'template.xlsx'
        self.client.post(reverse('upload_results'), {
            'session': self.session.id, 'term': self.term.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
            'excel_file': upload_file,
        })
        self.assertTrue(Result.objects.filter(student=self.active_student, subject=self.subject, session=self.session, term=self.term).exists())
        self.assertFalse(Result.objects.filter(student=self.inactive_student, subject=self.subject, session=self.session, term=self.term).exists())

    def test_include_inactive_applies_their_row(self):
        upload_file = self._build_workbook()
        upload_file.name = 'template.xlsx'
        self.client.post(reverse('upload_results'), {
            'session': self.session.id, 'term': self.term.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
            'excel_file': upload_file, 'include_inactive': 'on',
        })
        self.assertTrue(Result.objects.filter(student=self.inactive_student, subject=self.subject, session=self.session, term=self.term).exists())


class BulkMidtermResultUpdateActiveRosterTests(ActiveRosterFixtureMixin, TestCase):
    def test_default_excludes_inactive_student(self):
        response = self.client.post(reverse('bulk_update_midterm_results'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
        })
        students = [r.student for r in response.context['results']]
        self.assertIn(self.active_student, students)
        self.assertNotIn(self.inactive_student, students)
        self.assertFalse(MidTermResult.objects.filter(student=self.inactive_student).exists())

    def test_include_inactive_includes_them(self):
        response = self.client.post(reverse('bulk_update_midterm_results'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
            'include_inactive': 'on',
        })
        students = [r.student for r in response.context['results']]
        self.assertIn(self.inactive_student, students)


class DisplayMidtermResultsActiveRosterTests(ActiveRosterFixtureMixin, TestCase):
    def test_default_excludes_inactive_student(self):
        url = reverse('display_midterm_results', args=[self.session.id, self.term.id, self.school_class.id])
        response = self.client.get(url)
        students = [d['student'] for d in response.context['results_data']]
        self.assertIn(self.active_student, students)
        self.assertNotIn(self.inactive_student, students)

    def test_include_inactive_query_param_includes_them(self):
        url = reverse('display_midterm_results', args=[self.session.id, self.term.id, self.school_class.id])
        response = self.client.get(url, {'include_inactive': '1'})
        students = [d['student'] for d in response.context['results_data']]
        self.assertIn(self.inactive_student, students)


class ResultEntryTahfeezActiveRosterTests(ActiveRosterFixtureMixin, TestCase):
    def test_default_excludes_inactive_student(self):
        response = self.client.post(reverse('result_entry_tahfeez'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
        })
        students = list(response.context['students'])
        self.assertIn(self.active_student, students)
        self.assertNotIn(self.inactive_student, students)

    def test_include_inactive_includes_them(self):
        response = self.client.post(reverse('result_entry_tahfeez'), {
            'term': self.term.id, 'session': self.session.id,
            'class_assigned': self.school_class.id, 'subject': self.subject.id,
            'include_inactive': 'on',
        })
        students = list(response.context['students'])
        self.assertIn(self.inactive_student, students)
