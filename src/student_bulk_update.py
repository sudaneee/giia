"""
Bulk update of existing students' details, one class at a time.

Flow:
  1. Export  - pick a class, download an Excel file of its students with all
               their details.
  2. Edit    - change whatever is needed in Excel. Any cell left blank keeps
               the existing value, so nobody has to fill in every column.
  3. Preview - upload the file. Nothing is saved yet; the page lists every
               change (old -> new) and every row that couldn't be applied.
  4. Apply   - confirm the preview and the valid changes are saved in one
               transaction.

Rules:
  - Admission Number identifies the student and is never changed here.
  - Class is shown for reference only and is never changed here. A row whose
    student isn't currently in the selected class is rejected.
  - Students are only updated, never created (the existing bulk upload is for
    new students).
"""
from datetime import date, datetime

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from src.models import Guardian, SchoolClass, Student

SESSION_KEY = 'student_bulk_update_preview'

# (Excel header, key). Keys starting with "guardian_" map to the student's
# first guardian record; everything else maps to a Student field.
COLUMNS = [
    ('Admission Number', 'admission_number'),
    ('Class (reference only - not changed)', 'class'),
    ('First Name', 'first_name'),
    ('Last Name', 'last_name'),
    ('Date of Birth (YYYY-MM-DD)', 'date_of_birth'),
    ('Gender', 'gender'),
    ('Address', 'address'),
    ('Phone Number', 'phone_number'),
    ('Email', 'email'),
    ('Status (active/inactive/graduated/suspended)', 'status'),
    ('Admission Status (admitted/not_admitted)', 'admission_status'),
    ('Guardian First Name', 'guardian_first_name'),
    ('Guardian Last Name', 'guardian_last_name'),
    ('Guardian Phone Number', 'guardian_phone_number'),
    ('Guardian Email', 'guardian_email'),
    ('Guardian Relationship', 'guardian_relationship'),
]
HEADER_TO_KEY = {header: key for header, key in COLUMNS}
KEY_TO_LABEL = {key: header for header, key in COLUMNS}

STUDENT_STATUS_VALUES = {'active', 'inactive', 'graduated', 'suspended'}
ADMISSION_STATUS_VALUES = {'admitted', 'not_admitted'}

STUDENT_TEXT_FIELDS = ['first_name', 'last_name', 'gender', 'address', 'phone_number', 'email']
GUARDIAN_FIELDS = {
    'guardian_first_name': 'first_name',
    'guardian_last_name': 'last_name',
    'guardian_phone_number': 'phone_number',
    'guardian_email': 'email',
    'guardian_relationship': 'relationship',
}


def _text(value):
    """Excel cell -> stripped string. Whole-number floats lose the trailing .0."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _student_value_as_text(student, key):
    value = getattr(student, key)
    if key == 'date_of_birth':
        return value.isoformat() if value else ''
    return '' if value is None else str(value)


def _guardian_value_as_text(guardian, field):
    value = getattr(guardian, field, None) if guardian else None
    return '' if value is None else str(value)


# ----------------------------------------------------------------------
# Export
# ----------------------------------------------------------------------

@login_required(login_url='login')
def student_bulk_export(request):
    school_class = get_object_or_404(SchoolClass, id=request.GET.get('class_id'))
    students = (
        Student.objects.filter(enrolled_class=school_class)
        .prefetch_related('guardians')
        .order_by('last_name', 'first_name')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Students'
    ws.append([header for header, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)

    for student in students:
        guardian = student.guardians.first()
        row = []
        for _, key in COLUMNS:
            if key == 'class':
                row.append(str(school_class))
            elif key in GUARDIAN_FIELDS:
                row.append(_guardian_value_as_text(guardian, GUARDIAN_FIELDS[key]))
            else:
                row.append(_student_value_as_text(student, key))
        ws.append(row)

    for index, (header, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = max(14, min(40, len(header) + 2))

    filename = f"students_{school_class.name}_{school_class.arm}".replace(' ', '_') + '.xlsx'
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ----------------------------------------------------------------------
# Preview (parse + validate, saves nothing)
# ----------------------------------------------------------------------

def _parse_date(raw):
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = _text(raw)
    try:
        return datetime.strptime(text, '%Y-%m-%d').date()
    except ValueError:
        raise ValidationError('Date of birth must be YYYY-MM-DD.')


def _normalize_value(key, raw):
    """Returns the cleaned value to store, or raises ValidationError."""
    if key == 'date_of_birth':
        return _parse_date(raw).isoformat()
    text = _text(raw)
    if key == 'status':
        text = text.lower()
        if text not in STUDENT_STATUS_VALUES:
            raise ValidationError('Status must be active, inactive, graduated or suspended.')
    elif key == 'admission_status':
        text = text.lower().replace(' ', '_')
        if text not in ADMISSION_STATUS_VALUES:
            raise ValidationError('Admission status must be admitted or not_admitted.')
    elif key in ('email', 'guardian_email') and text:
        try:
            validate_email(text)
        except ValidationError:
            raise ValidationError(f'"{text}" is not a valid email address.')
    return text


def _build_preview(file_obj, school_class):
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.active

    header_row = [_text(c.value) for c in ws[1]]
    if 'Admission Number' not in header_row:
        return None, ['The file has no "Admission Number" column, so it can\'t be matched to students.']

    col_for_key = {}
    for col_index, header in enumerate(header_row):
        if header in HEADER_TO_KEY:
            col_for_key[HEADER_TO_KEY[header]] = col_index

    rows = []
    errors = []
    seen_admission_numbers = set()

    for row_number, excel_row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def cell(key):
            index = col_for_key.get(key)
            return excel_row[index] if index is not None and index < len(excel_row) else None

        if all(_text(v) == '' for v in excel_row):
            continue  # fully blank row

        admission_number = _text(cell('admission_number'))
        if not admission_number:
            errors.append(f'Row {row_number}: no admission number, skipped.')
            continue
        if admission_number in seen_admission_numbers:
            errors.append(f'Row {row_number}: {admission_number} appears more than once, skipped.')
            continue
        seen_admission_numbers.add(admission_number)

        student = Student.objects.filter(admission_number=admission_number).first()
        if not student:
            errors.append(f'Row {row_number}: no student with admission number {admission_number}.')
            continue
        if student.enrolled_class_id != school_class.id:
            errors.append(f'Row {row_number}: {admission_number} is not in {school_class}, skipped.')
            continue

        row_errors = []
        student_updates = {}
        guardian_updates = {}
        changes = []

        for key in list(STUDENT_TEXT_FIELDS) + ['date_of_birth', 'status', 'admission_status']:
            if key not in col_for_key:
                continue
            raw = cell(key)
            if _text(raw) == '' and not isinstance(raw, (date, datetime)):
                continue  # blank cell keeps the existing value
            try:
                new_value = _normalize_value(key, raw)
            except ValidationError as e:
                row_errors.append(f'{KEY_TO_LABEL[key]}: {e.messages[0]}')
                continue
            old_value = _student_value_as_text(student, key)
            if new_value != old_value:
                student_updates[key] = new_value
                changes.append({'field': KEY_TO_LABEL[key], 'old': old_value, 'new': new_value})

        guardian = student.guardians.first()
        for key, field in GUARDIAN_FIELDS.items():
            if key not in col_for_key:
                continue
            raw = cell(key)
            if _text(raw) == '':
                continue
            try:
                new_value = _normalize_value(key, raw)
            except ValidationError as e:
                row_errors.append(f'{KEY_TO_LABEL[key]}: {e.messages[0]}')
                continue
            old_value = _guardian_value_as_text(guardian, field)
            if new_value != old_value:
                guardian_updates[field] = new_value
                changes.append({'field': KEY_TO_LABEL[key], 'old': old_value, 'new': new_value})

        if guardian is None and guardian_updates:
            if not (guardian_updates.get('first_name') and guardian_updates.get('last_name')):
                row_errors.append('Guardian: a new guardian needs at least a first and last name.')
                guardian_updates = {}
                changes = [c for c in changes if not c['field'].startswith('Guardian')]

        if row_errors:
            errors.append(f'Row {row_number} ({admission_number}): ' + '; '.join(row_errors) + ' - this row was not applied.')
            continue
        if not changes:
            continue

        rows.append({
            'student_id': student.id,
            'admission_number': admission_number,
            'name': f'{student.first_name} {student.last_name}',
            'changes': changes,
            'student_updates': student_updates,
            'guardian_updates': guardian_updates,
            'needs_new_guardian': guardian is None and bool(guardian_updates),
        })

    return {'class_id': school_class.id, 'class_name': str(school_class), 'rows': rows}, errors


# ----------------------------------------------------------------------
# Apply (only what the preview stored; re-checks each student first)
# ----------------------------------------------------------------------

def _apply_preview(preview):
    applied = 0
    skipped = []
    with transaction.atomic():
        for item in preview['rows']:
            student = Student.objects.filter(id=item['student_id']).first()
            if not student or student.enrolled_class_id != preview['class_id']:
                skipped.append(item['admission_number'])
                continue

            for key, value in item['student_updates'].items():
                if key == 'date_of_birth':
                    value = date.fromisoformat(value)
                setattr(student, key, value)
            student.save()

            if item['guardian_updates']:
                guardian = student.guardians.first()
                if guardian is None:
                    guardian = Guardian.objects.create(relationship=item['guardian_updates'].get('relationship', ''),
                                                       **{k: v for k, v in item['guardian_updates'].items() if k != 'relationship'})
                    student.guardians.add(guardian)
                else:
                    for field, value in item['guardian_updates'].items():
                        setattr(guardian, field, value)
                    guardian.save()
            applied += 1
    return applied, skipped


# ----------------------------------------------------------------------
# Page
# ----------------------------------------------------------------------

@login_required(login_url='login')
def student_bulk_update(request):
    school_classes = SchoolClass.objects.all()
    url = reverse('student_bulk_update')

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'cancel':
            request.session.pop(SESSION_KEY, None)
            request.session.pop('student_bulk_update_errors', None)
            messages.info(request, 'Bulk update cancelled. Nothing was changed.')
            return redirect(url)

        if action == 'apply':
            preview = request.session.pop(SESSION_KEY, None)
            request.session.pop('student_bulk_update_errors', None)
            if not preview:
                messages.error(request, 'There is no preview to apply. Upload the file again.')
                return redirect(url)
            applied, skipped = _apply_preview(preview)
            message = f'Updated {applied} student(s) in {preview["class_name"]}.'
            if skipped:
                message += f' Skipped {len(skipped)} because they changed since the preview: {", ".join(skipped)}.'
            messages.success(request, message)
            return redirect(url)

        # Otherwise: an uploaded file to preview
        school_class = get_object_or_404(SchoolClass, id=request.POST.get('class_id'))
        excel_file = request.FILES.get('excel_file')
        if not excel_file or not excel_file.name.endswith('.xlsx'):
            messages.error(request, 'Please upload an Excel file (.xlsx).')
            return redirect(url)

        preview, errors = _build_preview(excel_file, school_class)
        if preview is None:
            messages.error(request, errors[0])
            return redirect(url)

        request.session[SESSION_KEY] = preview
        request.session['student_bulk_update_errors'] = errors
        return redirect(url)

    preview = request.session.get(SESSION_KEY)
    errors = request.session.get('student_bulk_update_errors', [])
    selected_class_id = (preview or {}).get('class_id') or request.GET.get('class_id')

    return render(request, 'src/student_bulk_update.html', {
        'school_classes': school_classes,
        'selected_class_id': str(selected_class_id) if selected_class_id else '',
        'preview': preview,
        'errors': errors,
        'applicable_count': len(preview['rows']) if preview else 0,
    })
